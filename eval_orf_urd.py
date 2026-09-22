# -*- coding: utf-8 -*-
"""
STT evaluation for EGRA ORF_URD (Oral Reading Fluency, Urdu connected text).

Scope: ORF_URD only. Deepgram nova-3 (language=ur) transcripts, sliced to each
ORF_URD window using word-level timestamps, aligned against the 60-word gold
passage, then compared to the enumerator's per-word correct/incorrect taps.

Two questions are answered separately:
  Q1  How well did the ASR capture the target passage?   -> WER / CER
  Q2  Could the ASR replace the human scorer?            -> per-item agreement,
      i.e. does an ASR mismatch predict a human "incorrect" tap?  P/R/F1, CPM
"""
import csv, json, re, unicodedata, pathlib, collections

ROOT = pathlib.Path(__file__).parent
TASK = "ORF_URD"
import os
PAD = float(os.environ.get("PAD", 5.0))  # slack around the manifest window.
# Keep this SMALL: the reading comprehension Q&A follows immediately after and the
# enumerator repeats passage words ("bilal ne kya ... dekha"), so a wide window lets
# the aligner match the enumerator's questions instead of the child's read.
OUT = ROOT / "eval_orf_urd"
OUT.mkdir(exist_ok=True)

# ---------------------------------------------------------------- normalisation
_DIAC = dict.fromkeys(range(0x064B, 0x0653))        # arabic diacritics
_DIAC[0x0670] = None                                 # superscript alef
_DIAC[0x0640] = None                                 # tatweel
for _z in (0x200B, 0x200C, 0x200D):                  # zero-width chars
    _DIAC[_z] = None
_MAP = {"ي": "ی", "ى": "ی", "ك": "ک",
        "ه": "ہ", "ۀ": "ہ", "أ": "ا",
        "إ": "ا", "آ": "ا", "ٱ": "ا",
        "ؤ": "و", "ئ": "ی"}
_PUNCT = re.compile(u"[۔،؟!\\.\\,\\?\\:\\;\"'\\(\\)\\[\\]«»\\-–—/]")


def norm(w):
    w = unicodedata.normalize("NFC", str(w)).translate(_DIAC)
    w = "".join(_MAP.get(c, c) for c in w)
    return _PUNCT.sub("", w).strip()


def norm_seq(words):
    return [n for n in (norm(w) for w in words) if n]


# ---------------------------------------------------------------- alignment
def align(ref, hyp):
    """Levenshtein alignment; returns list of (op, ref_i, hyp_j)."""
    n, m = len(ref), len(hyp)
    d = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n + 1):
        d[i][0] = i
    for j in range(m + 1):
        d[0][j] = j
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            c = 0 if ref[i - 1] == hyp[j - 1] else 1
            d[i][j] = min(d[i - 1][j] + 1, d[i][j - 1] + 1, d[i - 1][j - 1] + c)
    ops, i, j = [], n, m
    while i > 0 or j > 0:
        if i > 0 and j > 0 and d[i][j] == d[i - 1][j - 1] + (0 if ref[i - 1] == hyp[j - 1] else 1):
            ops.append(("eq" if ref[i - 1] == hyp[j - 1] else "sub", i - 1, j - 1))
            i, j = i - 1, j - 1
        elif i > 0 and d[i][j] == d[i - 1][j] + 1:
            ops.append(("del", i - 1, None))
            i -= 1
        else:
            ops.append(("ins", None, j - 1))
            j -= 1
    return ops[::-1]


def align_infix(ref, hyp):
    """Align ref as an INFIX of hyp: hyp prefix/suffix outside the match is free.

    The ORF window runs start_orf_urd -> start_rdcomp_urd, so it brackets the
    child's reading with the enumerator's instructions on both sides. Charging
    those as insertions inflates WER past 1.0, so they are not scored; only the
    best-matching span inside the window is.
    """
    n, m = len(ref), len(hyp)
    d = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n + 1):
        d[i][0] = i
    for j in range(m + 1):
        d[0][j] = 0                      # free start anywhere in hyp
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            c = 0 if ref[i - 1] == hyp[j - 1] else 1
            d[i][j] = min(d[i - 1][j] + 1, d[i][j - 1] + 1, d[i - 1][j - 1] + c)
    j = min(range(m + 1), key=lambda jj: d[n][jj])   # free end anywhere in hyp
    end = j
    ops, i = [], n
    while i > 0:
        if j > 0 and d[i][j] == d[i - 1][j - 1] + (0 if ref[i - 1] == hyp[j - 1] else 1):
            ops.append(("eq" if ref[i - 1] == hyp[j - 1] else "sub", i - 1, j - 1))
            i, j = i - 1, j - 1
        elif d[i][j] == d[i - 1][j] + 1:
            ops.append(("del", i - 1, None))
            i -= 1
        else:
            ops.append(("ins", None, j - 1))
            j -= 1
    return ops[::-1], j, end


def wer(ref, hyp):
    if not ref:
        return None, (0, 0, 0)
    ops = align_infix(ref, hyp)[0]
    s = sum(1 for o in ops if o[0] == "sub")
    dl = sum(1 for o in ops if o[0] == "del")
    ins = sum(1 for o in ops if o[0] == "ins")
    return (s + dl + ins) / len(ref), (s, dl, ins)


def cer(ref, hyp):
    r, h = list(" ".join(ref)), list(" ".join(hyp))
    if not r:
        return None
    return wer(r, h)[0]


# ---------------------------------------------------------------- gold
import openpyxl

wb = openpyxl.load_workbook(ROOT / "EGRA_STT_eval_key.xlsx", read_only=True)
GOLD = [r[5] for r in wb["Item_Key"].iter_rows(values_only=True) if r[0] == TASK]
GOLD_N = norm_seq(GOLD)
PASSAGE = next(r[6] for r in wb["Passages"].iter_rows(values_only=True) if r[0] == TASK)
assert len(GOLD) == 60, len(GOLD)

human = collections.defaultdict(dict)
with open(ROOT / "gold_items_long.csv", encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        if r["task_id"] == TASK:
            human[r["KEY"]][int(float(r["item_index"]))] = int(r["is_correct"])

man = {}
with open(ROOT / "eval_manifest.csv", encoding="utf-8-sig") as f:
    for r in csv.DictReader(f):
        if r["task_id"] == TASK:
            man[r["KEY"]] = r


# ---------------------------------------------------------------- transcripts
def words_for(stem):
    j = json.loads((ROOT / "transcripts" / "json" / (stem + ".json")).read_text(encoding="utf-8"))
    return j["results"]["channels"][0]["alternatives"][0]["words"]


# Join on the FULL 36-char uuid, not a prefix. Filenames are AA_<uuid>_enumerator.m4a,
# which is exactly the `file=` parameter of the manifest's audio_url column.
_UUID = re.compile(r"AA_([0-9a-fA-F-]{36})_enumerator")
STEMS = {}
for _p in (ROOT / "transcripts" / "json").glob("*.json"):
    _m = _UUID.search(_p.stem)
    if not _m:
        raise SystemExit("cannot parse uuid from transcript filename: %s" % _p.name)
    STEMS["uuid:" + _m.group(1).lower()] = _p.stem

# ---------------------------------------------------------------- evaluate
summary, per_item_rows = [], []
for key, stem in sorted(STEMS.items()):
    short = key[5:13]                      # display only
    if key not in man:
        print("no %s manifest row for %s" % (TASK, key))
        continue
    m = man[key]
    t0, t1 = float(m["offset_start_sec"]), float(m["offset_end_sec"])
    attempted = int(float(m["attempted"]))

    # pad the window: the read can start late or run past start_rdcomp_urd.
    # infix alignment ignores whatever extra we pull in.
    ws = [w for w in words_for(stem) if t0 - PAD <= w["start"] < t1 + PAD]
    hyp = norm_seq(w.get("punctuated_word", w["word"]) for w in ws)

    # the child only reached `attempted` words -> reference is the truncated passage
    ref = GOLD_N[:attempted]
    w_rate, (sub, dele, ins) = wer(ref, hyp)
    c_rate = cer(ref, hyp)

    # ---- per-item: did the ASR recover gold item i? -> predicted "read correctly"
    ops, mstart, mend = align_infix(ref, hyp)
    pred = {i: 0 for i in range(1, attempted + 1)}
    for op, ri, hj in ops:
        if op == "eq" and ri is not None:
            pred[ri + 1] = 1
    hum = human.get(key, {})

    tp = fp = fn = tn = 0  # positive class = "incorrect" (what the enumerator taps)
    for i in range(1, attempted + 1):
        h = hum.get(i)
        p = pred[i]
        if h is None:
            continue
        per_item_rows.append({"uuid": short, "item_index": i, "gold_word": GOLD[i - 1],
                              "human_correct": h, "asr_recovered": p,
                              "agree": int(h == p)})
        if h == 0 and p == 0:
            tp += 1
        elif h == 1 and p == 0:
            fp += 1
        elif h == 0 and p == 1:
            fn += 1
        else:
            tn += 1

    n = tp + fp + fn + tn
    prec = tp / (tp + fp) if tp + fp else None
    rec = tp / (tp + fn) if tp + fn else None
    f1 = 2 * prec * rec / (prec + rec) if prec and rec else None

    # duration of the matched reading span only, not the whole padded window
    mw = ws[mstart:mend]
    span = (mw[-1]["end"] - mw[0]["start"]) if mw else 0
    mconf = (sum(w["confidence"] for w in mw) / len(mw)) if mw else None
    match_t0 = round(mw[0]["start"], 1) if mw else None
    match_t1 = round(mw[-1]["end"], 1) if mw else None
    hum_corr = sum(hum.get(i, 0) for i in range(1, attempted + 1))
    asr_corr = sum(pred.values())
    dur = float(m["segment_len_sec"])
    summary.append({
        "uuid": short, "attempted": attempted, "seg_start": t0, "seg_end": t1,
        "seg_len_sec": dur, "asr_words_in_window": len(hyp), "asr_span_sec": round(span, 1),
        "match_start": match_t0, "match_end": match_t1,
        "asr_mean_conf_in_read": round(mconf, 3) if mconf is not None else None,
        "wer": round(w_rate, 4) if w_rate is not None else None,
        "cer": round(c_rate, 4) if c_rate is not None else None,
        "sub": sub, "del": dele, "ins": ins,
        "human_correct": hum_corr, "asr_correct": asr_corr,
        "correct_delta": asr_corr - hum_corr,
        "human_cpm": round(hum_corr / span * 60, 1) if span else None,
        "asr_cpm": round(asr_corr / span * 60, 1) if span else None,
        "item_agreement": round((tp + tn) / n, 4) if n else None,
        "precision_incorrect": round(prec, 4) if prec is not None else None,
        "recall_incorrect": round(rec, 4) if rec is not None else None,
        "f1_incorrect": round(f1, 4) if f1 is not None else None,
        "tp": tp, "fp": fp, "fn": fn, "tn": tn,
        "asr_text": " ".join(hyp[mstart:mend]),
    })

# ---------------------------------------------------------------- write
with open(OUT / "summary.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(summary[0].keys()))
    w.writeheader()
    w.writerows(summary)
with open(OUT / "per_item.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(per_item_rows[0].keys()))
    w.writeheader()
    w.writerows(per_item_rows)

TOT = lambda k: sum(r[k] for r in summary)
tp, fp, fn, tn = TOT("tp"), TOT("fp"), TOT("fn"), TOT("tn")
P = tp / (tp + fp) if tp + fp else 0
R = tp / (tp + fn) if tp + fn else 0
print("EGRA ORF_URD - STT evaluation (Deepgram nova-3, language=ur)")
print("passage: %s\n" % PASSAGE[:80])
hdr = "%-9s %4s %5s %5s %6s %6s %6s %7s %7s"
print(hdr % ("uuid", "att", "WER", "CER", "hum_c", "asr_c", "agree", "P(inc)", "R(inc)"))
for r in summary:
    fmt = lambda v: ("%.2f" % v) if v is not None else "-"
    print(hdr % (r["uuid"], r["attempted"], fmt(r["wer"]), fmt(r["cer"]),
                 r["human_correct"], r["asr_correct"], fmt(r["item_agreement"]),
                 fmt(r["precision_incorrect"]), fmt(r["recall_incorrect"])))
print("\nPOOLED over %d students / %d items" % (len(summary), tp + fp + fn + tn))
print("  item agreement       %.3f" % ((tp + tn) / (tp + fp + fn + tn)))
print("  precision(incorrect) %.3f   recall(incorrect) %.3f   F1 %.3f"
      % (P, R, 2 * P * R / (P + R) if P + R else 0))
print("  confusion: tp=%d fp=%d fn=%d tn=%d" % (tp, fp, fn, tn))
n_all = tp + fp + fn + tn
n_inc = tp + fn
bP = n_inc / n_all
bF1 = 2 * bP * 1.0 / (bP + 1.0)
print("\nBASELINE - label every word 'incorrect' without listening to anything:")
print("  precision %.3f   recall 1.000   F1 %.3f   (base rate of incorrect = %.1f%%)"
      % (bP, bF1, 100 * bP))
print("  -> the ASR scorer beats this trivial baseline by %.3f F1" % ((2*P*R/(P+R) if P+R else 0) - bF1))
with open(OUT / "side_by_side.txt", "w", encoding="utf-8") as f:
    f.write("GOLD PASSAGE (60 words)\n%s\n\n" % PASSAGE)
    for r in summary:
        f.write("=" * 70 + "\n")
        f.write("%s  attempted=%d  human_correct=%d  asr_recovered=%d\n"
                % (r["uuid"], r["attempted"], r["human_correct"], r["asr_correct"]))
        f.write("matched span %s-%ss  mean_conf=%s  WER=%.2f\n"
                % (r["match_start"], r["match_end"], r["asr_mean_conf_in_read"], r["wer"]))
        f.write("REF: %s\n" % " ".join(GOLD_N[:r["attempted"]]))
        f.write("ASR: %s\n\n" % r["asr_text"])

print("\n-> eval_orf_urd/summary.csv, per_item.csv, side_by_side.txt")
